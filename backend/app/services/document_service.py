from collections.abc import Sequence
from dataclasses import dataclass
from datetime import date

from fastapi import UploadFile
from pathlib import Path
from docx import Document as DocxDocument
from pypdf import PdfReader, PasswordType
from sqlalchemy import delete, func, select
from sqlalchemy.orm import Session

from app.Schemas import DocumentOut, SearchHit
from app.config import get_settings
from app.models import DocumentChunk, DocumentStatus, PolicyDocument
from app.services.bm25_store import (
    get_bm25_index,
    invalidate_bm25_index,
    search_children_bm25,
)
from app.services.chunking import split_policy_text
from app.services.fusion import fusion_candidate_k, rrf_fuse
# 名字里的 _hits 是为了不和 search_documents 的 rerank 参数撞名
from app.services.reranker import rerank as rerank_hits
from app.services.milvus_store import (
    NO_PAGE,
    date_from_int,
    delete_document,
    drop_collection,
    effective_date_to_int,
    get_milvus_client,
    search_children,
    upsert_child_chunks,
)
from app.services.vector_store import get_embeddings

ALLOWED_EXTENSIONS ={".pdf",".docx",".txt"}

async def save_upload(session:Session,
        upload:UploadFile,
        topic:str | None = None,
        effective_date:date | None = None,):
    settings=get_settings()
    suffix = Path(upload.filename or "").suffix.lower()
    if suffix not in ALLOWED_EXTENSIONS:
        raise ValueError("仅支持PDF、DOCX、TXT文件")
    content=await upload.read()
    if len(content) > settings. max_upload_mb*1024*1024:
        raise ValueError("文件大小超出限制")

    #上传目录来自配置，不存在就创建（exist_ok=True 本身就是幂等的）
    upload_path=settings.resolved_upload_path
    upload_path.mkdir(parents=True,exist_ok=True)

    document=PolicyDocument(
        filename=upload.filename or f"policy{suffix}",
        stored_path="",
        content_type=upload.content_type,
        topic=topic,
        effective_date=effective_date,
        #还没有做切片和向量化，所以状态是 pending 而不是 ready
        status=DocumentStatus.pending
    )
    session.add(document)

    #先flush拿到数据库生成的id，再用id给文件命名，避免同名文件互相覆盖
    session.flush()

    uploadPath = upload_path / f"{document.id}{suffix}"
    uploadPath.write_bytes(content)
    document.stored_path=str(uploadPath)
    session.commit()

    #created_at/updated_at 是数据库的 server_default，要refresh才能取回真实值
    session.refresh(document)
    return DocumentOut.model_validate(document)

def ingest_upload(
        session:Session,
        upload:UploadFile,
        topic:str | None = None,
        effective_date:date | None = None,
):
    suffix = Path(upload.filename or "").suffix.lower()
    settings=get_settings()
    if suffix not in ALLOWED_EXTENSIONS:
        raise ValueError("仅支持PDF、DOCX、TXT文件")
    #同步读而不是 await upload.read()：调用方（uploads 路由）是普通 def，
    #整个函数在线程池里跑，用底层的同步文件对象直接读即可
    content=upload.file.read()
    if len(content) > settings.max_upload_mb*1024*1024:
        raise ValueError("文件大小超出限制")
    upload_path=settings.resolved_upload_path
    upload_path.mkdir(parents=True,exist_ok=True)
    document = PolicyDocument(
        filename=upload.filename or f"policy{suffix}",
        stored_path="",
        content_type=upload.content_type,
        topic=topic,
        effective_date=effective_date,
        status=DocumentStatus.indexing
    )
    session.add(document)

    # 先flush拿到数据库生成的id，再用id给文件命名，避免同名文件互相覆盖
    session.flush()

    uploadPath = upload_path / f"{document.id}{suffix}"
    uploadPath.write_bytes(content)
    document.stored_path = str(uploadPath)

    #切片 + 落库 + 向量化这一段和纯文本入库完全一样，抽成 helper 共用
    _chunk_persist_and_index(session, document, extract_pages(uploadPath))

    # created_at/updated_at 是数据库的 server_default，要refresh才能取回真实值
    session.refresh(document)
    return DocumentOut.model_validate(document)


def ingest_text(
        session:Session,
        *,
        title:str | None,
        content:str,
        topic:str | None = None,
        effective_date:date | None = None,
):
    """纯文本入库：没有文件，正文直接走和上传文件一模一样的切片 + 向量化流程。"""
    if not content or not content.strip():
        raise ValueError("content 不能为空")
    document = PolicyDocument(
        filename=title or "text-input",
        #没有磁盘文件，删除时靠 stored_path 为空来跳过删文件这一步
        stored_path="",
        #content_type 列不允许为空，纯文本固定标成 text/plain
        content_type="text/plain",
        topic=topic,
        effective_date=effective_date,
        status=DocumentStatus.indexing,
    )
    session.add(document)

    #先flush拿到数据库生成的id，DocumentChunk.document_id 才能挂上去
    session.flush()

    #整篇算一页，页码为 None，和 docx/txt 的 extract_pages 口径一致
    _chunk_persist_and_index(session, document, [(None, content)])

    session.refresh(document)
    return DocumentOut.model_validate(document)


def _chunk_persist_and_index(
        session:Session,
        document:PolicyDocument,
        pages:list[tuple[int | None,str]],
) -> None:
    """把 (页码, 文本) 列表切片、落 MySQL、向量化写 Milvus，并把 document 置为 ready。

    调用前 document 必须已经 session.add + flush（拿到 id）。
    保持原有顺序：先落块再写 Milvus，最坏情况只是「MySQL 有块、Milvus 没有」，
    重跑一次即可幂等修复。
    """
    #父子切块：父块和子块都进 MySQL，只有子块随后向量化写入 Milvus
    chunk_rows:list[DocumentChunk] = []
    # 第 i 个子块的父块在 chunk_rows 里的下标。
    # chunk.id 要 flush 之后才有值，所以先记下标，flush 完再回填 parent_id
    parent_pos:dict[int,int] = {}
    ordinal = 0
    #此处的page是页码，text是页内容
    for page,text in pages:
        # 切片结果的结构是「父块紧跟它自己的子块」，所以最近一个父块就是子块的父块
        current_parent_pos:int | None = None
        for draft in split_policy_text(text,page=page):
            chunk=DocumentChunk(
                document_id=document.id,
                ordinal=ordinal,
                section=draft.section,
                content=draft.content,
                is_parent=draft.is_parent,
                # document_chunks 表没有 page 列，页码放进已有的 JSON 字段
                metadata_json={**draft.metadata, "page": draft.page},
            )
            session.add(chunk)
            chunk_rows.append(chunk)
            if draft.is_parent:
                current_parent_pos = len(chunk_rows) - 1
            elif current_parent_pos is not None:
                parent_pos[len(chunk_rows) - 1] = current_parent_pos
            ordinal += 1

    # flush 之后每个 chunk 才有 id，父块的 id 才能回填给子块的 parent_id
    session.flush()
    for child_index, parent_index in parent_pos.items():
        chunk_rows[child_index].parent_id = chunk_rows[parent_index].id

    child_rows = [chunk for chunk in chunk_rows if not chunk.is_parent]
    document.chunk_count = len(child_rows)
    # 先把切块落库。Milvus 写入无法和 MySQL 组成分布式事务，
    # 这个顺序保证最坏情况只是「MySQL 有块、Milvus 没有」，重跑一次即可幂等修复
    session.commit()

    try:
        if child_rows:
            vectors = get_embeddings().embed_documents(
                [chunk.content or "" for chunk in child_rows]
            )
            upsert_child_chunks(
                get_milvus_client(),
                document.id,
                [
                    {
                        "chunk_id": chunk.id,
                        "vector": vector,
                        "document_id": document.id,
                        "parent_id": chunk.parent_id or "",
                        "content": chunk.content or "",
                        "section": chunk.section or "",
                        # Milvus 标量字段不支持 null，没有页码时用 -1 占位
                        "page": chunk.metadata_json.get("page") if chunk.metadata_json.get("page") is not None else NO_PAGE,
                        "ordinal": chunk.ordinal,
                        # 主题和生效日期从文档上冗余下来，检索时才能在 Milvus 里原生过滤
                        "topic": document.topic or "",
                        "effective_date": effective_date_to_int(document.effective_date),
                    }
                    for chunk, vector in zip(child_rows, vectors)
                ],
            )
    except Exception as exc:
        # Milvus 侧失败：标成 failed 并清掉可能写了一半的向量。
        # MySQL 里的切块故意保留，方便排查问题
        document.status = DocumentStatus.failed
        document.error_message = f"向量化写入 Milvus 失败：{exc}"
        session.commit()
        try:
            delete_document(get_milvus_client(), document.id)
        except Exception:
            pass
        raise

    document.status=DocumentStatus.ready
    document.index_version = 1
    document.error_message = None
    session.commit()

    #BM25 语料变了（新文档从 indexing 变 ready 才会进语料），让缓存索引作废。
    #必须在 commit 之后：提前失效的话，落在失效与 commit 之间的检索会把
    #「还没提交」的语料建进索引，而那份缓存再也没人来失效它
    invalidate_bm25_index()

def list_documents(session:Session,
        topic:str | None = None,
        status:DocumentStatus | None = None,
        page:int = 1,
        size:int = 20,) -> tuple[Sequence[PolicyDocument], int]:
    """按条件分页查询文档，返回(当前页数据, 符合条件的总条数)"""
    #筛选条件只构造一次，count 和分页查询共用，
    #这样 total 和 items 的口径一定是一致的
    conditions=[]
    if topic:
        #topic 列上有索引，用精确匹配才能命中；LIKE '%x%' 会退化成全表扫描
        conditions.append(PolicyDocument.topic == topic)
    if status is not None:
        conditions.append(PolicyDocument.status == status)

    total=session.scalar(
        select(func.count()).select_from(PolicyDocument).where(*conditions)
    ) or 0

    statement=(
        select(PolicyDocument)
        .where(*conditions)
        #created_at 只精确到秒，同一秒入库的多条记录顺序不稳定，
        #翻页时会出现重复或漏数据，所以用 id 做第二排序键
        .order_by(PolicyDocument.created_at.desc(),PolicyDocument.id.desc())
        .offset((page-1)*size)
        .limit(size)
    )
    return session.scalars(statement).all(), total

def extract_pages(path:Path) ->list[tuple[int | None,str]]:
    """把文件拆成 (页码, 文本) 列表：pdf 逐页，docx/txt 整篇算一页（页码为 None）。"""
    suffix = Path(path).suffix.lower()
    pages:list[tuple[int | None,str]] = []
    if suffix == ".pdf":
        for index,page in enumerate(PdfReader(path).pages,start=1):
            # extract_text() 对纯图片页会返回 None
            pages.append((index, page.extract_text() or ""))
    elif suffix == ".docx":
        document = DocxDocument(str(path))
        content = "\n".join(paragraph.text for paragraph in document.paragraphs)
        pages.append((None, content))
    else:
        content = Path(path).read_text(encoding="utf-8")
        pages.append((None, content))
    return pages


#检索结果的状态。抽成常量是为了让路由和文档引用同一份字面量，不用手抄字符串
STATUS_OK = "ok"
STATUS_NO_EVIDENCE = "no_evidence"
STATUS_REFUSED = "refused"
REASON_NO_EVIDENCE = "检索未命中任何证据"
REASON_REFUSED = "召回证据与问题相关性不足，已拒答"


@dataclass(frozen=True)
class SearchOutcome:
    """检索结果 + 拒答状态。

    status 不能从 list[SearchHit] 反推（no_evidence 和 refused 都是空 hits，
    但前者是「库里没有」、后者是「有但质量不够」，上层处理方式不同），所以单独带出来。
    """
    hits: list[SearchHit]
    status: str
    reason: str | None


def _build_milvus_filter(
        document_id:str | None,
        topic:str | None,
        effective_before:date | None,
) -> str:
    """三条过滤条件用 and 串成 Milvus 表达式；空表达式是合法的，Milvus 当成「无过滤」。

    过滤下推给 Milvus 参与检索，而不是检索完再回 MySQL 后过滤——
    后者会丢召回：top_k=8 里被筛掉几条就只剩几条，库里明明还有符合条件的也不会返回。
    """
    conditions:list[str] = []
    if document_id:
        conditions.append(_filter_literal("document_id", document_id))
    if topic:
        conditions.append(_filter_literal("topic", topic))
    if effective_before:
        #没填生效日期的文档存的是 DATE_MAX，会被 <= 自然排除在外。
        #BM25 那边用同一个 effective_date_to_int，两路口径才对得上
        conditions.append(f"effective_date <= {effective_date_to_int(effective_before)}")
    return " and ".join(conditions)


def _materialize_hits(
        session:Session,
        ordered_ids:list[str],
        dense_by_id:dict[str,dict],
        bm25_by_id:dict[str,dict],
        rrf_by_id:dict[str,float],
        rerank_by_id:dict[str,float],
) -> list[SearchHit]:
    """把排好序的 chunk_id 列表还原成 SearchHit（含 small-to-big 的父块正文）。

    dense 和 bm25 的两份 dict 内容同源（都是 MySQL 的同一行 + 文档的 topic/生效日期），
    所以标量字段谁在就用谁当 base；distance 只从 dense 取、bm25_score 只从 bm25 取，
    rerank_score 只从 rerank 旁表取——这样「这条到底靠哪一路命中、重排给了多少分」
    在响应里都是可分辨的，重排不会把原始分数抹掉。
    """
    #small-to-big：子块负责被检索命中，父块负责给 LLM 足够上下文，
    #所以按 parent_id 一次性把父块正文查回来（对最终命中的父块取并集，只查一次）
    parent_ids = set()
    for chunk_id in ordered_ids:
        base = dense_by_id.get(chunk_id) or bm25_by_id.get(chunk_id)
        if base and base.get("parent_id"):
            parent_ids.add(base["parent_id"])

    parent_text:dict[str,str] = {}
    if parent_ids:
        rows = session.scalars(
            select(DocumentChunk).where(DocumentChunk.id.in_(parent_ids))
        ).all()
        parent_text = {row.id: (row.content or "") for row in rows}

    hits:list[SearchHit] = []
    for chunk_id in ordered_ids:
        dense = dense_by_id.get(chunk_id)
        bm25 = bm25_by_id.get(chunk_id)
        base = dense or bm25
        if base is None:
            #理论不可达：ordered_ids 全部来自这两份 dict
            continue
        parent_id = base["parent_id"]
        hits.append(SearchHit(
            chunk_id=chunk_id,
            document_id=base["document_id"],
            #Milvus 标量字段不支持 null，入库时没有父块写的是空串，这里还原成 None
            parent_id=parent_id or None,
            content=base["content"],
            section=base["section"] or None,
            #没有页码时两路存的都是 NO_PAGE(-1)，这里同样还原成 None
            page=None if base["page"] in (None, NO_PAGE) else base["page"],
            ordinal=base["ordinal"],
            #COSINE 度量下 distance 就是相似度，越高越相关；BM25-only 命中没有它
            distance=dense["distance"] if dense else None,
            bm25_score=bm25["bm25_score"] if bm25 else None,
            rrf_score=rrf_by_id.get(chunk_id),
            rerank_score=rerank_by_id.get(chunk_id),
            #回显主题和生效日期，方便肉眼确认过滤条件真的生效了
            topic=base["topic"] or None,
            effective_date=date_from_int(base["effective_date"]),
            parent_content=parent_text.get(parent_id) if parent_id else None,
        ))
    return hits


def search_documents(
        session:Session,
        query:str,
        top_k:int = 8,
        document_id:str | None = None,
        topic:str | None = None,
        effective_before:date | None = None,
        mode:str = "hybrid",
        rerank:bool | None = None,
) -> SearchOutcome:
    """检索子块，可选重排，再按 parent_id 回 MySQL 取父块正文（small-to-big）。

    mode="vector"：纯向量（过滤表达式下推给 Milvus）
    mode="bm25"  ：纯关键词，只走内存 BM25，不调 embedding
    mode="hybrid"：两路各取候选后用 RRF 融合（默认）

    rerank：cross-encoder 重排阶段，叠加在三种 mode 之上、不是新的 mode。
    传 None 就跟随 settings.enable_reranker。开启时：候选池缩到 pool_k（CPU 成本），
    逐对打分并 sigmoid 归一化，丢掉低于 min_evidence_score 的，再取 final_k 条。
    一条都不剩就是拒答（status=refused），这正是「证据不足就别硬答」的落点。

    两条路的过滤条件语义必须逐条对齐（见 bm25_store._matches），
    否则同一个查询在不同 mode 下返回的命中集合会不一致。
    """
    settings = get_settings()
    use_rerank = settings.enable_reranker if rerank is None else rerank
    #最终条数：重排开启时还要被 RERANK_TOP_K 压一道（它表达的是「喂给 LLM 的证据条数」）
    final_k = min(top_k, settings.rerank_top_k) if use_rerank else top_k
    #候选池：RETRIEVAL_TOP_K 读起来就是「喂给重排模型的候选数」，所以以它为基础；
    #max(..., final_k) 保证 top_k 比它大时候选池够用；外层的 fusion_candidate_k
    #保持今天的天花板不被突破。重排是对 (query, 全文) 逐对跑大模型，
    #池子不封顶的话 32 对在 CPU 上要好几秒，封到 8 对是这里最主要的动机
    pool_k = (
        min(max(settings.retrieval_top_k, final_k), fusion_candidate_k(top_k))
        if use_rerank else 0
    )

    dense_hits:list[dict] = []
    bm25_hits:list[dict] = []

    if mode in ("vector","hybrid"):
        #重排开启时一律取 pool_k；否则维持现状（vector 模式必须保持 limit=top_k 原样）
        limit = pool_k if use_rerank else (top_k if mode == "vector" else fusion_candidate_k(top_k))
        #embed_query 会自动带上 BGE 的检索指令前缀（见 vector_store.py），文档侧不加
        vector = get_embeddings().embed_query(query)
        dense_hits = search_children(
            get_milvus_client(), vector, top_k=limit,
            filter_expr=_build_milvus_filter(document_id, topic, effective_before),
        )

    if mode in ("bm25","hybrid"):
        limit = pool_k if use_rerank else (top_k if mode == "bm25" else fusion_candidate_k(top_k))
        #只在真要走这条路的模式下才装载索引（纯向量请求不该为 BM25 付首次构建的代价）
        bm25_hits = search_children_bm25(
            get_bm25_index(session), query, top_k=limit,
            document_id=document_id, topic=topic, effective_before=effective_before,
        )

    dense_by_id = {hit["chunk_id"]: hit for hit in dense_hits}
    bm25_by_id = {hit["chunk_id"]: hit for hit in bm25_hits}
    rrf_by_id:dict[str,float] = {}
    rerank_by_id:dict[str,float] = {}

    if mode == "hybrid":
        #dense 放在 rankings[0]：RRF 同分时「先出现的排前面」，让向量路在并列时占先
        fused = rrf_fuse([
            [hit["chunk_id"] for hit in dense_hits],
            [hit["chunk_id"] for hit in bm25_hits],
        ])
        rrf_by_id = dict(fused)
        ranked = [chunk_id for chunk_id,_ in fused]
    else:
        #单路模式：各自已经排好序
        source = dense_hits if mode == "vector" else bm25_hits
        ranked = [hit["chunk_id"] for hit in source]

    if not use_rerank:
        ordered_ids = ranked[:top_k]
    else:
        #重排阶段：夹在「有序 id + 旁表构造」与 _materialize_hits 之间，
        #也就是在 small-to-big 那一次父块查询之前，不多一次 DB 往返
        candidates = [
            base
            for base in (dense_by_id.get(chunk_id) or bm25_by_id.get(chunk_id)
                         for chunk_id in ranked[:pool_k])
            if base is not None
        ]
        if not candidates:
            #召回本身就是空的：不加载模型，直接判「没有证据」
            return SearchOutcome([], STATUS_NO_EVIDENCE, REASON_NO_EVIDENCE)
        kept = rerank_hits(
            query, candidates,
            top_k=final_k, min_score=settings.min_evidence_score,
        )
        if not kept:
            #召回到了，但全被阈值滤掉 —— 这才是「拒答」
            return SearchOutcome([], STATUS_REFUSED, REASON_REFUSED)
        rerank_by_id = dict(kept)
        ordered_ids = [chunk_id for chunk_id,_ in kept]

    hits = _materialize_hits(
        session, ordered_ids, dense_by_id, bm25_by_id, rrf_by_id, rerank_by_id
    )
    if not hits:
        #兜底：ordered_ids 非空却一条都还原不出来（理论不可达）
        return SearchOutcome([], STATUS_NO_EVIDENCE, REASON_NO_EVIDENCE)
    return SearchOutcome(hits, STATUS_OK, None)


def _filter_literal(field:str, value:str) -> str:
    """把用户输入安全地拼进 Milvus 过滤表达式。

    Milvus 表达式没有参数绑定，pymilvus 客户端也完全不校验、原样转发给服务端，
    所以引号和反斜杠必须自己挡掉，否则可以闭合字符串往表达式里塞条件
    """
    if '"' in value or "\\" in value:
        raise ValueError(f"{field} 含非法字符（引号或反斜杠）")
    return f'{field} == "{value}"'


def delete_one(session:Session, document_id:str) -> bool:
    """删除单个文档：Milvus 向量、MySQL 记录、磁盘文件一起清掉。文档不存在返回 False。

    顺序刻意和摄入时相反（摄入是「先 MySQL 后 Milvus」）：这里**先删 Milvus**，
    失败就整个请求失败、MySQL 行还在，重试即可；若反过来先删了 MySQL 再挂，
    向量就成了 API 再也够不着的孤儿。
    """
    document = session.get(PolicyDocument, document_id)
    if document is None:
        return False
    stored_path = document.stored_path

    #delete_document 自带 has_collection 守卫，重复删除是幂等的
    delete_document(get_milvus_client(), document_id)

    #document_chunks.parent_id 是自引用外键：父行若先于子行被删，MySQL 会报 1451。
    #所以先删子块再删父块——这个顺序无论外键是 CASCADE 还是 RESTRICT 都成立。
    for is_parent in (False, True):
        session.execute(
            delete(DocumentChunk)
            .where(
                DocumentChunk.document_id == document_id,
                DocumentChunk.is_parent.is_(is_parent),
            )
            #批量 delete 不走 ORM 身份映射，关掉同步，免得它去核对马上要丢弃的对象
            .execution_options(synchronize_session=False)
        )
    session.execute(
        delete(PolicyDocument)
        .where(PolicyDocument.id == document_id)
        .execution_options(synchronize_session=False)
    )
    session.commit()

    #被删文档的子块已从 MySQL 消失，BM25 语料变了，缓存索引作废
    invalidate_bm25_index()

    #磁盘文件最后删、失败不阻断：留个孤儿文件无害，重试也不会更糟。
    #必须先 is_file()：纯文本入库的 stored_path 是空串，而 Path("") 会塌缩成当前目录
    if stored_path:
        try:
            path = Path(stored_path)
            if path.is_file():
                path.unlink()
        except OSError:
            pass
    return True


def reset_all(session:Session) -> dict[str,int]:
    """清空全部数据但保留表结构：Milvus collection、MySQL 两张表、上传目录里的文件。"""
    #直接 drop collection 而不是按条件 delete：后者只写墓碑，segment 和 HNSW 索引都还在，
    #语义上不算「清空」。下次入库时 ensure_collection 会自动按 schema 重建，无需任何补救代码
    drop_collection(get_milvus_client())

    #批量 delete 不触发 ORM 的 cascade，所以顺序是硬要求：
    #chunks 必须先于 documents，且块内部先子后父（自引用外键）
    chunks_deleted = 0
    for is_parent in (False, True):
        chunks_deleted += session.execute(
            delete(DocumentChunk)
            .where(DocumentChunk.is_parent.is_(is_parent))
            .execution_options(synchronize_session=False)
        ).rowcount
    documents_deleted = session.execute(
        delete(PolicyDocument).execution_options(synchronize_session=False)
    ).rowcount
    session.commit()

    #全库清空，BM25 语料整个没了，缓存索引作废
    invalidate_bm25_index()

    #上传目录本身保留，只清里面的文件
    files_deleted = 0
    upload_path = get_settings().resolved_upload_path
    if upload_path.is_dir():
        for child in upload_path.iterdir():
            try:
                if child.is_file():
                    child.unlink()
                    files_deleted += 1
            except OSError:
                #个别文件删不掉不影响整体重置，跳过
                continue

    return {
        "documents_deleted": documents_deleted,
        "chunks_deleted": chunks_deleted,
        "files_deleted": files_deleted,
    }